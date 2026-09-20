"""Content extraction for all supported file types (PDF, TXT, Markdown, DOCX, HTML,
images), including Phase 4 multimodal processing: OCR fallback for scanned PDFs and
images, and structured table extraction for PDF/DOCX.

Returns an `ExtractionResult` — a unified representation the chunker consumes
regardless of source format. HTML and DOCX body text is normalized into the same
heading-marked plain text the Markdown chunker already understands (headings become
`#`..`######` lines) rather than teaching the chunker format-specific structure.
Tables are pulled out as separate `ExtractedTable` objects instead — per the project
brief, "do not treat a table as plain text if structured extraction is possible" — and
turned into their own chunks by the ingestion pipeline (see
app/services/ingestion/pipeline.py), not flattened into surrounding body text.

Images: if OCR (Tesseract, via app/services/extraction/ocr.py) finds real text, the
image behaves like any other text document from here on (`is_visual_only=False`). If
OCR finds nothing — or Tesseract isn't installed — the image stays `is_visual_only=True`
and the ingestion pipeline tries a vision-LLM caption instead; if neither is available,
it's cataloged (format/dimensions) without being searchable, same as Phase 2.
"""
from __future__ import annotations

from dataclasses import dataclass, field

from pypdf import PdfReader

from app.core.config import Settings, get_settings
from app.services.extraction.ocr import ocr_image_bytes, ocr_pdf_pages


@dataclass
class ExtractedPage:
    page_number: int
    text: str
    section: str | None = None


@dataclass
class ExtractedTable:
    page: int
    headers: list[str]
    rows: list[list[str]]


@dataclass
class ExtractionResult:
    pages: list[ExtractedPage]
    page_count: int
    metadata: dict = field(default_factory=dict)
    is_visual_only: bool = False
    tables: list[ExtractedTable] = field(default_factory=list)


class UnsupportedFileTypeError(ValueError):
    pass


class CorruptedFileError(ValueError):
    pass


def render_table_markdown(table: ExtractedTable) -> str:
    """Renders a structured table as a Markdown table — the text form used for
    embedding/search. The structured `headers`/`rows` (see ExtractedTable) remain the
    source of truth, stored verbatim in the chunk's metadata; this is only the
    searchable serialization of it."""
    lines = ["| " + " | ".join(table.headers) + " |"]
    lines.append("|" + "|".join(["---"] * len(table.headers)) + "|")
    for row in table.rows:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def extract(file_type: str, raw_bytes: bytes, settings: Settings | None = None) -> ExtractionResult:
    settings = settings or get_settings()
    if file_type == "pdf":
        return _extract_pdf(raw_bytes, settings)
    if file_type == "txt":
        return _extract_plain_text(raw_bytes)
    if file_type == "markdown":
        return _extract_markdown(raw_bytes)
    if file_type == "docx":
        return _extract_docx(raw_bytes)
    if file_type == "html":
        return _extract_html(raw_bytes)
    if file_type == "image":
        return _extract_image(raw_bytes, settings)
    raise UnsupportedFileTypeError(f"No extractor registered for file_type={file_type!r}")


def _decode(raw_bytes: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return raw_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise CorruptedFileError("Could not decode file as text")


def _extract_pdf_tables(raw_bytes: bytes) -> list[ExtractedTable]:
    import io

    import pdfplumber

    tables: list[ExtractedTable] = []
    try:
        with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
            for page_number, page in enumerate(pdf.pages, start=1):
                for raw_table in page.extract_tables() or []:
                    if len(raw_table) < 2:
                        continue  # need at least a header row + one data row
                    header = [(cell or "").strip() for cell in raw_table[0]]
                    rows = [[(cell or "").strip() for cell in row] for row in raw_table[1:]]
                    if not any(header) or not any(any(r) for r in rows):
                        continue
                    tables.append(ExtractedTable(page=page_number, headers=header, rows=rows))
    except Exception:
        # Table detection is a best-effort enhancement; a pdfplumber failure must not
        # fail extraction when pypdf already produced usable text.
        return []
    return tables


def _extract_pdf(raw_bytes: bytes, settings: Settings) -> ExtractionResult:
    import io

    try:
        reader = PdfReader(io.BytesIO(raw_bytes))
    except Exception as exc:  # pypdf raises various errors for malformed PDFs
        raise CorruptedFileError(f"Could not parse PDF: {exc}") from exc

    if len(reader.pages) == 0:
        raise CorruptedFileError("PDF has no pages")

    pages = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        pages.append(ExtractedPage(page_number=i, text=text))

    doc_info = reader.metadata or {}
    metadata: dict = {"title": getattr(doc_info, "title", None)}

    if not any(p.text.strip() for p in pages):
        # No text layer at all -> likely a scanned PDF. Try OCR before giving up.
        ocr_pages = ocr_pdf_pages(raw_bytes, settings)
        if ocr_pages and any(p.strip() for p in ocr_pages):
            pages = [
                ExtractedPage(page_number=i, text=text)
                for i, text in enumerate(ocr_pages, start=1)
            ]
            metadata["ocr_used"] = True
        else:
            raise CorruptedFileError(
                "PDF has no extractable text layer and OCR found nothing (or is unavailable)"
            )

    tables = _extract_pdf_tables(raw_bytes)
    return ExtractionResult(pages=pages, page_count=len(reader.pages), metadata=metadata, tables=tables)


def _extract_plain_text(raw_bytes: bytes) -> ExtractionResult:
    text = _decode(raw_bytes)
    if not text.strip():
        raise CorruptedFileError("File is empty")
    return ExtractionResult(pages=[ExtractedPage(page_number=1, text=text)], page_count=1)


def _extract_markdown(raw_bytes: bytes) -> ExtractionResult:
    text = _decode(raw_bytes)
    if not text.strip():
        raise CorruptedFileError("File is empty")
    # Markdown is treated as a single logical page; section headings are recovered
    # during chunking (structure-aware chunker splits on `#`/`##`).
    return ExtractionResult(pages=[ExtractedPage(page_number=1, text=text)], page_count=1)


_MAX_DOCX_UNCOMPRESSED_BYTES = 300 * 1024 * 1024  # 300 MB
_MAX_DOCX_COMPRESSION_RATIO = 100  # uncompressed:compressed


def _reject_if_zip_bomb(raw_bytes: bytes) -> None:
    """Production hardening: a .docx file is a zip archive, and neither `zipfile`
    nor `python-docx` guards against a decompression bomb — a small, crafted upload
    that expands to gigabytes in memory when opened, a real denial-of-service risk
    that `MAX_UPLOAD_SIZE_BYTES` alone doesn't stop (it only caps the *compressed*
    upload size). Inspects each entry's *declared* size from the zip's own central
    directory — cheap, and doesn't require decompressing anything — and rejects
    before `DocxDocument()` ever touches the content if the total declared
    uncompressed size, or any single entry's compression ratio, looks like a bomb
    rather than an ordinary document."""
    import io
    import zipfile

    try:
        with zipfile.ZipFile(io.BytesIO(raw_bytes)) as zf:
            total_uncompressed = 0
            for info in zf.infolist():
                total_uncompressed += info.file_size
                if info.compress_size > 0 and info.file_size / info.compress_size > _MAX_DOCX_COMPRESSION_RATIO:
                    raise CorruptedFileError("DOCX rejected: a compressed entry's expansion ratio is too high")
                if total_uncompressed > _MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise CorruptedFileError("DOCX rejected: declared uncompressed size exceeds the safety limit")
    except zipfile.BadZipFile:
        # Not a valid zip at all — DocxDocument() below will raise its own,
        # more specific CorruptedFileError for this; nothing further to check here.
        return


def _extract_docx(raw_bytes: bytes) -> ExtractionResult:
    import io
    import zipfile

    from docx import Document as DocxDocument
    from docx.opc.exceptions import PackageNotFoundError

    _reject_if_zip_bomb(raw_bytes)

    try:
        docx_doc = DocxDocument(io.BytesIO(raw_bytes))
    except (PackageNotFoundError, KeyError, ValueError, zipfile.BadZipFile) as exc:
        raise CorruptedFileError(f"Could not parse DOCX: {exc}") from exc

    lines: list[str] = []
    for paragraph in docx_doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = (paragraph.style.name if paragraph.style else "") or ""
        if style_name.startswith("Heading"):
            level = "".join(ch for ch in style_name if ch.isdigit()) or "1"
            level = min(int(level), 6)
            lines.append(f"{'#' * level} {text}")
        elif style_name == "Title":
            lines.append(f"# {text}")
        else:
            lines.append(text)

    # Tables get structured extraction (headers + rows), not flattened into body text
    # — see ExtractedTable / render_table_markdown and the module docstring.
    tables: list[ExtractedTable] = []
    for table in docx_doc.tables:
        rows_text = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        rows_text = [r for r in rows_text if any(r)]
        if len(rows_text) < 2:
            continue
        tables.append(ExtractedTable(page=1, headers=rows_text[0], rows=rows_text[1:]))

    full_text = "\n\n".join(lines)
    if not full_text.strip() and not tables:
        raise CorruptedFileError("DOCX has no extractable text content")

    props = docx_doc.core_properties
    metadata = {"title": props.title or None, "author": props.author or None}
    return ExtractionResult(
        pages=[ExtractedPage(page_number=1, text=full_text)] if full_text.strip() else [],
        page_count=1,
        metadata=metadata,
        tables=tables,
    )


def _extract_html(raw_bytes: bytes) -> ExtractionResult:
    from bs4 import BeautifulSoup

    html_text = _decode(raw_bytes)
    soup = BeautifulSoup(html_text, "html.parser")

    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()

    title = soup.title.string.strip() if soup.title and soup.title.string else None

    body = soup.body or soup
    lines: list[str] = []
    heading_tags = {f"h{i}" for i in range(1, 7)}
    for element in body.find_all(list(heading_tags | {"p", "li", "td", "th"})):
        text = " ".join(element.get_text(" ", strip=True).split())
        if not text:
            continue
        if element.name in heading_tags:
            level = int(element.name[1])
            lines.append(f"{'#' * level} {text}")
        else:
            lines.append(text)

    full_text = "\n\n".join(lines)
    if not full_text.strip():
        raise CorruptedFileError("HTML has no extractable text content")

    return ExtractionResult(
        pages=[ExtractedPage(page_number=1, text=full_text)], page_count=1, metadata={"title": title}
    )


def _extract_image(raw_bytes: bytes, settings: Settings) -> ExtractionResult:
    import io

    from PIL import Image, UnidentifiedImageError

    try:
        with Image.open(io.BytesIO(raw_bytes)) as img:
            width, height = img.size
            image_format = img.format
    except UnidentifiedImageError as exc:
        raise CorruptedFileError(f"Could not parse image: {exc}") from exc

    ocr_text = ocr_image_bytes(raw_bytes, settings)
    base_metadata = {
        "image_width": width,
        "image_height": height,
        "image_format": image_format,
    }

    if len(ocr_text) >= settings.min_ocr_text_length:
        return ExtractionResult(
            pages=[ExtractedPage(page_number=1, text=ocr_text)],
            page_count=1,
            metadata={**base_metadata, "ocr_used": True, "text_extraction": "ocr"},
            is_visual_only=False,
        )

    return ExtractionResult(
        pages=[],
        page_count=1,
        metadata={**base_metadata, "ocr_used": False, "text_extraction": "pending_visual_description"},
        is_visual_only=True,
    )
